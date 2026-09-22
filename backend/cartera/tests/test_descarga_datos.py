"""Descarga del archivo que alimenta un dashboard, con todas sus filas y columnas.

Lo que importa verificar es que sea el conjunto COMPLETO —no el recorte que muestra cada
sección— y que salga saneado: los valores vienen de una base que este proyecto no controla, y un
texto que empiece con `=` abierto en Excel se ejecuta como fórmula.
"""

import datetime
import uuid
from unittest.mock import patch

import pandas as pd
from django.conf import settings
from django.contrib.auth import get_user_model
from django.test import TestCase
from openpyxl import load_workbook
from rest_framework.test import APIClient

from cartera.exceptions import CarteraError
from cartera.models import CargaArchivo
from cartera.services import descarga_datos
from cartera.services.carga_archivos import HOJA_ARCHIVO_PERMANENTE
from cartera.services.dashboards import crear_dashboard
from cartera.utils.archivos import asegurar_directorio

User = get_user_model()

DASHBOARD = 'finanzas'
CORTE = datetime.date(2026, 9, 20)


def _df():
    return pd.DataFrame({
        'Cliente': ['ACME', 'BETA', '=CMD("formato")'],
        'Saldo Total': [1000.5, 200.0, 0.0],
        'Fecha de Vencimiento': [
            datetime.date(2026, 9, 30), datetime.date(2026, 8, 1), datetime.date(2026, 7, 1),
        ],
        'Zona': ['Norte', 'Sur', None],
    })


def _carga(df=None, fecha_corte=CORTE):
    df = _df() if df is None else df
    carga = CargaArchivo.objects.create(
        id=uuid.uuid4(), dashboard_id=DASHBOARD, nombre_original='origen.xlsx',
        estado=CargaArchivo.Estado.PROCESADO, fecha_corte=fecha_corte,
        total_filas_excel=len(df), filas_validas=len(df),
    )
    nombre = f'{carga.id}.xlsx'
    ruta = asegurar_directorio(settings.CARTERA_ARCHIVOS_DIR) / nombre
    df.to_excel(ruta, index=False, sheet_name=HOJA_ARCHIVO_PERMANENTE)
    carga.archivo_permanente_nombre = nombre
    carga.save(update_fields=['archivo_permanente_nombre'])
    return carga


def _hoja(buffer):
    libro = load_workbook(buffer)
    return libro[descarga_datos.HOJA]


class ExcelDeLaCargaVigenteTests(TestCase):
    def setUp(self):
        crear_dashboard(nombre='Finanzas')
        _carga()

    def test_trae_todas_las_columnas_y_todas_las_filas(self):
        buffer, _nombre = descarga_datos.excel_de_la_carga_vigente(DASHBOARD)
        hoja = _hoja(buffer)

        encabezados = [celda.value for celda in hoja[1]]
        self.assertEqual(encabezados, ['Cliente', 'Saldo Total', 'Fecha de Vencimiento', 'Zona'])
        # Una fila por registro más la de encabezados.
        self.assertEqual(hoja.max_row, len(_df()) + 1)

    def test_los_valores_llegan_como_valores_y_no_como_texto(self):
        buffer, _nombre = descarga_datos.excel_de_la_carga_vigente(DASHBOARD)
        hoja = _hoja(buffer)

        self.assertEqual(hoja.cell(row=2, column=2).value, 1000.5)
        self.assertEqual(hoja.cell(row=2, column=3).value.date(), datetime.date(2026, 9, 30))

    def test_un_valor_que_empieza_con_igual_no_queda_como_formula(self):
        """Viene de una base que este proyecto no controla: sin el saneo, Excel lo ejecutaría.

        El dataframe se inyecta parcheando el lector en vez de escribirlo al .xlsx de la carga: al
        guardarlo ahí, openpyxl lo convierte en fórmula y al releerlo vuelve vacío, así que la
        prueba estaría comprobando el archivo intermedio y no el saneo de la descarga.
        """
        con_formula = pd.DataFrame({'Cliente': ['ACME', '=CMD("formato")'], 'Saldo Total': [1.0, 2.0]})
        with patch.object(descarga_datos.carga_archivos, 'leer_archivo_de_carga',
                          return_value=('ruta.xlsx', con_formula)):
            buffer, _nombre = descarga_datos.excel_de_la_carga_vigente(DASHBOARD)

        sospechoso = _hoja(buffer).cell(row=3, column=1).value
        self.assertTrue(sospechoso.startswith("'="), sospechoso)

    def test_una_celda_vacia_queda_vacia_y_no_dice_nan(self):
        buffer, _nombre = descarga_datos.excel_de_la_carga_vigente(DASHBOARD)
        hoja = _hoja(buffer)
        # La tercera fila de datos trae la zona vacía.
        self.assertIn(hoja.cell(row=4, column=4).value, ('', None))

    def test_el_nombre_del_archivo_identifica_el_dashboard_y_el_corte(self):
        _buffer, nombre = descarga_datos.excel_de_la_carga_vigente(DASHBOARD)
        self.assertEqual(nombre, 'Finanzas - datos 2026-09-20.xlsx')

    def test_sin_fecha_de_corte_el_nombre_lo_dice_en_vez_de_quedar_a_medias(self):
        CargaArchivo.objects.filter(dashboard_id=DASHBOARD).update(fecha_corte=None)
        _buffer, nombre = descarga_datos.excel_de_la_carga_vigente(DASHBOARD)
        self.assertIn('sin-corte', nombre)

    def test_descarga_la_carga_vigente_y_no_una_anterior(self):
        # Lo que se descarga tiene que coincidir con lo que la pantalla está mostrando.
        _carga(df=pd.DataFrame({'Cliente': ['NUEVO'], 'Saldo Total': [9.0]}))
        buffer, _nombre = descarga_datos.excel_de_la_carga_vigente(DASHBOARD)
        hoja = _hoja(buffer)
        self.assertEqual([c.value for c in hoja[1]], ['Cliente', 'Saldo Total'])
        self.assertEqual(hoja.cell(row=2, column=1).value, 'NUEVO')


class SinDatosTests(TestCase):
    def test_un_dashboard_sin_cargas_lo_dice_con_un_mensaje_claro(self):
        crear_dashboard(nombre='Finanzas')
        with self.assertRaises(CarteraError) as ctx:
            descarga_datos.excel_de_la_carga_vigente(DASHBOARD)
        self.assertEqual(ctx.exception.codigo, 'SIN_DATOS_PARA_DESCARGAR')


class EndpointTests(TestCase):
    def setUp(self):
        crear_dashboard(nombre='Finanzas')
        _carga()
        self.client = APIClient()
        self.client.force_authenticate(
            user=User.objects.create_superuser(username='u', email='u@example.com', password='Clave-Segura-123'),
        )

    def test_devuelve_un_xlsx_como_adjunto(self):
        resp = self.client.get(f'/api/dashboards/{DASHBOARD}/descargar-datos')

        self.assertEqual(resp.status_code, 200)
        self.assertEqual(
            resp['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        self.assertIn('attachment;', resp['Content-Disposition'])
        self.assertIn('Finanzas - datos 2026-09-20.xlsx', resp['Content-Disposition'])

    def test_sin_sesion_no_se_descarga(self):
        resp = APIClient().get(f'/api/dashboards/{DASHBOARD}/descargar-datos')
        self.assertEqual(resp.status_code, 401)

    def test_sin_datos_responde_el_error_de_negocio_de_siempre(self):
        CargaArchivo.objects.filter(dashboard_id=DASHBOARD).delete()
        resp = self.client.get(f'/api/dashboards/{DASHBOARD}/descargar-datos')
        self.assertEqual(resp.status_code, 400)
        self.assertEqual(resp.json()['error'], 'SIN_DATOS_PARA_DESCARGAR')
